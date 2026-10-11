"""Super platform libraries, scoped RAG/memory, real tools and MCP."""
import asyncio
from contextlib import asynccontextmanager
from datetime import timedelta
import hashlib
import io
import ipaddress
import json
import math
from pathlib import Path
import re
import secrets
import socket
import statistics
import time
from typing import Literal
from urllib.parse import urlsplit
import uuid
from datetime import datetime

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import Response
import httpx
from pydantic import BaseModel, ConfigDict, Field, field_validator

from .platform_store import KINDS, Resource, store, timestamp, validate_resource

router = APIRouter(prefix="/api/platform", tags=["super-platform"])
RUNTIMES = {}


class SearchRequest(BaseModel):
    query: str = Field(default="", max_length=4000)
    knowledge_ids: list[str] | None = Field(default=None, max_length=100)
    memory_ids: list[str] | None = Field(default=None, max_length=100)
    scope: str | None = Field(default=None, max_length=100)
    limit: int = Field(default=5, ge=1, le=20)


class ToolRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    arguments: dict = Field(default_factory=dict)
    token: str | None = Field(default=None, max_length=200)
    knowledge_ids: list[str] = Field(default_factory=list, max_length=100)
    memory_scope: str = Field(default="global", min_length=1, max_length=100)
    project_id: str | None = None
    snapshot_id: str | None = None
    mcp_server_ids: list[str] = Field(default_factory=list, max_length=20)


class FlowRunRequest(BaseModel):
    prompt: str = Field(min_length=1, max_length=12000)
    knowledge_ids: list[str] = Field(default_factory=list, max_length=100)
    memory_ids: list[str] = Field(default_factory=list, max_length=100)
    memory_scope: str = Field(default="global", min_length=1, max_length=100)
    remember: bool = False
    project_id: str | None = None
    snapshot_id: str | None = None
    mcp_server_ids: list[str] = Field(default_factory=list, max_length=20)


def schema(properties, required=()):
    return {"type": "object", "properties": properties, "required": list(required), "additionalProperties": False}


STRING = {"type": "string"}
TOOLS = [
    {"id": "math", "name": "数学计算", "description": "有限数值的四则运算；拒绝非有限结果。", "parameters": schema({"a": {"type": "number"}, "b": {"type": "number"}, "operation": {"type": "string", "enum": ["add", "subtract", "multiply", "divide"]}}, ["a", "b"])},
    {"id": "text_stats", "name": "文本统计", "description": "统计字符、单词、行数。", "parameters": schema({"text": STRING}, ["text"])},
    {"id": "statistics", "name": "数值统计", "description": "均值、中位数、标准差、极值与计数。", "parameters": schema({"values": {"type": "array", "items": {"type": "number"}, "maxItems": 10000}}, ["values"])},
    {"id": "markdown_to_text", "name": "Markdown 转文本", "description": "转换 Markdown 内容为纯文本。", "parameters": schema({"text": STRING}, ["text"])},
    {"id": "web_fetch", "name": "网页正文", "description": "读取公开 HTTP(S) 网页；返回 URL 和提取正文。", "parameters": schema({"url": STRING}, ["url"])},
    {"id": "web_search", "name": "联网搜索", "description": "通过 DDGS 搜索公开网页，返回标题、链接和摘要。", "parameters": schema({"query": STRING, "max_results": {"type": "integer", "minimum": 1, "maximum": 10}}, ["query"])},
    {"id": "knowledge_search", "name": "知识库检索", "description": "对明确选择的文档进行 BM25 检索，返回文档和行范围。", "parameters": schema({"query": STRING, "limit": {"type": "integer", "minimum": 1, "maximum": 20}}, ["query"])},
    {"id": "memory_read", "name": "读取记忆", "description": "读取当前记忆空间中的相关条目。", "parameters": schema({"query": STRING, "limit": {"type": "integer", "minimum": 1, "maximum": 20}})},
    {"id": "memory_save", "name": "保存记忆", "description": "保存到当前记忆空间，并关联原任务。", "parameters": schema({"name": STRING, "content": STRING}, ["content"])},
    {"id": "project_read", "name": "读取项目快照", "description": "只读取已选择快照中的路径与原始行范围。", "parameters": schema({"path": STRING, "start_line": {"type": "integer", "minimum": 1}, "end_line": {"type": "integer", "minimum": 1}}, ["path"])},
    {"id": "mcp_call", "name": "MCP 工具", "description": "调用已选择服务器中明确允许的工具。", "parameters": schema({"server_id": STRING, "tool": STRING, "arguments": {"type": "object"}}, ["server_id", "tool"])},
]


def clean_ids(values):
    try:
        return list(dict.fromkeys(str(uuid.UUID(value)) for value in values))
    except (ValueError, TypeError) as error:
        raise HTTPException(422, "资源 ID 无效") from error


def validate_context(value):
    from .projects import store as projects
    for kind, field in [("knowledge", "knowledge_ids"), ("memories", "memory_ids"), ("mcp", "mcp_server_ids")]:
        for identifier in clean_ids(value.get(field, [])):
            store.get(kind, identifier)
    if bool(value.get("project_id")) != bool(value.get("snapshot_id")):
        raise HTTPException(422, "项目上下文需要同时选择项目与快照")
    if value.get("project_id"):
        projects.snapshot(value["project_id"], value["snapshot_id"], True)


def prepare_runtime(job):
    """Capability token exists only in process memory and the private worker input."""
    context = {key: job.get(key) for key in ["knowledge_ids", "memory_ids", "memory_scope", "project_id", "snapshot_id", "mcp_server_ids"]}
    context.update(job_id=job["id"], tools=set())
    for agent in (job.get("workflow") or {}).get("agents", []):
        context["tools"].update(agent.get("tools", []))
    # Text tools for ordinary chat/group must be explicitly enabled as well.
    context["tools"].update(job.get("tools", []))
    token = secrets.token_urlsafe(32)
    RUNTIMES[token] = {**context, "expires_at": time.monotonic() + 1800}
    return {"token": token, **{key: value for key, value in context.items() if key != "tools"}}


def runtime_context(token):
    context = RUNTIMES.get(token or "")
    if not context or context["expires_at"] < time.monotonic():
        raise HTTPException(403, "任务工具授权不存在或已到期")
    return context


def release_runtime(token):
    RUNTIMES.pop(token, None)


def context_material(job):
    from .projects import store as projects
    pieces, references = [], []
    if job.get("knowledge_ids"):
        found = store.search("knowledge", job["prompt"], job["knowledge_ids"], limit=6)
        for item in found:
            pieces.append(f"[知识文档 {item['name']}，第 {item['start_line']}–{item['end_line']} 行]\n{item['content']}")
            references.append({"kind": "knowledge", **item})
    if job.get("memory_ids"):
        for identifier in job["memory_ids"]:
            item = store.get("memories", identifier)
            pieces.append(f"[用户选择的记忆 {item['name']}]\n{item['content'][:3000]}")
            references.append({"kind": "memory", "id": item["id"], "name": item["name"], "scope": item["scope"]})
    if isinstance(job.get("workflow"), dict) and job["workflow"].get("memory"):
        found = store.search("memories", job["prompt"], scope=job.get("memory_scope") or "global", limit=5)
        for item in found:
            pieces.append(f"[记忆空间 {item['scope']}：{item['name']}]\n{item['content']}")
            references.append({"kind": "memory", **item})
    if job.get("project_id") and job["mode"] != "review":
        snapshot = projects.snapshot(job["project_id"], job["snapshot_id"], True)
        for item in snapshot["files"]:
            current = item.get("current") or item.get("base")
            if current:
                pieces.append(f"[项目 {snapshot['project_name']}：{item['path']}，第 {current['start_line']}–{current['end_line']} 行]\n{current['text']}")
        references.append({"kind": "project", "project_id": job["project_id"], "snapshot_id": snapshot["id"], "digest": snapshot["digest"], "coverage": snapshot["coverage"]})
    content = "\n\n".join(pieces)
    truncated = len(content) > 8000
    return content[:8000], references, truncated


def remember_result(job):
    if not job.get("remember"):
        return None
    text = json.dumps(job.get("result"), ensure_ascii=False)[:12000]
    return store.put("memories", {"name": job["prompt"][:100], "content": "任务：" + job["prompt"][:2000] + "\n结果：" + text,
                                  "scope": job.get("memory_scope") or "global", "project_id": job.get("project_id"), "tags": [job["mode"]]},
                     metadata={"job_id": job["id"], "models": job["models"], "source": "explicit_remember"})


async def public_url(value):
    try:
        url = urlsplit(value)
        if url.scheme not in {"http", "https"} or not url.hostname or url.username or url.password or url.port not in {None, 80, 443}:
            raise ValueError()
        addresses = await asyncio.to_thread(socket.getaddrinfo, url.hostname, url.port or 443, type=socket.SOCK_STREAM)
        if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
            raise ValueError()
    except (ValueError, OSError) as error:
        raise HTTPException(422, "网页工具只接受不含凭据的公开 HTTP(S) 地址") from error
    return value


async def web_fetch(url):
    from trafilatura import extract
    for _ in range(4):
        await public_url(url)
        async with httpx.AsyncClient(timeout=20, follow_redirects=False, trust_env=True) as client:
            async with client.stream("GET", url, headers={"User-Agent": "ModelDesk/2.0 research-reader"}) as response:
                if response.is_redirect:
                    from urllib.parse import urljoin
                    url = urljoin(url, response.headers.get("location", ""))
                    continue
                response.raise_for_status()
                raw = bytearray()
                async for data in response.aiter_bytes():
                    raw.extend(data)
                    if len(raw) > 2_000_000:
                        raise HTTPException(422, "网页超过 2 MB 上限")
                text = raw.decode(response.encoding or "utf-8", errors="replace")
                content = extract(text, include_links=True, output_format="txt") or re.sub(r"<[^>]+>", " ", text)
                return {"url": url, "content": content[:16000], "truncated": len(content) > 16000}
    raise HTTPException(422, "网页重定向次数过多")


async def web_search(query, max_results=5):
    from ddgs import DDGS
    def search():
        results = list(DDGS(timeout=15).text(query, max_results=max_results, backend="auto"))
        return [{"title": item.get("title", ""), "url": item.get("href", ""), "content": item.get("body", "")[:2000]} for item in results]
    results = await asyncio.wait_for(asyncio.to_thread(search), timeout=25)
    if not results:
        raise HTTPException(503, "联网搜索没有返回结果，请换关键词或检查网络")
    return results


@asynccontextmanager
async def mcp_session(config):
    from mcp import ClientSession
    if config["transport"] == "sse":
        from mcp.client.sse import sse_client
        context = sse_client(config["url"], timeout=15, sse_read_timeout=20)
    else:
        from mcp.client.streamable_http import streamablehttp_client
        context = streamablehttp_client(config["url"], timeout=timedelta(seconds=15), sse_read_timeout=timedelta(seconds=20))
    async with context as streams:
        async with ClientSession(streams[0], streams[1], read_timeout_seconds=timedelta(seconds=20)) as session:
            await session.initialize()
            yield session


async def mcp_discover(identifier):
    config = store.get("mcp", identifier)
    if not config["enabled"]:
        raise HTTPException(409, "MCP 连接已停用")
    async with asyncio.timeout(25), mcp_session(config) as session:
        result = await session.list_tools()
        tools = [{"name": item.name, "description": item.description or "", "inputSchema": item.inputSchema,
                  "allowed": item.name in config["allowed_tools"]} for item in result.tools]
    return {"id": identifier, "online": True, "tools": tools}


async def mcp_call(identifier, name, arguments, selected=None):
    if selected is not None and identifier not in selected:
        raise HTTPException(403, "此任务未选择该 MCP 服务器")
    config = store.get("mcp", identifier)
    if not config["enabled"] or name not in config["allowed_tools"]:
        raise HTTPException(403, "该 MCP 工具未明确授权")
    async with asyncio.timeout(30), mcp_session(config) as session:
        result = await session.call_tool(name, arguments)
    if result.isError:
        raise HTTPException(502, "MCP 工具返回错误；请查看服务器状态")
    value = result.model_dump(mode="json")
    if len(json.dumps(value)) > 32000:
        return {"content": json.dumps(value, ensure_ascii=False)[:16000], "truncated": True}
    return value


async def invoke_tool(name, arguments, context):
    from jsonschema import validate, ValidationError
    definition = next((item for item in TOOLS if item["id"] == name), None)
    if not definition:
        raise HTTPException(404, "工具未注册")
    if len(json.dumps(arguments, ensure_ascii=False)) > 100000:
        raise HTTPException(422, "工具输入过大")
    try:
        validate(arguments, definition["parameters"])
    except ValidationError as error:
        raise HTTPException(422, "工具参数与声明的类型或范围不一致") from error
    if name == "math":
        a, b = float(arguments["a"]), float(arguments["b"])
        try:
            operation = arguments.get("operation", "add")
            result = {"add": lambda: a + b, "subtract": lambda: a - b, "multiply": lambda: a * b, "divide": lambda: a / b}[operation]()
            if not all(math.isfinite(value) for value in [a, b, result]):
                raise ValueError()
        except (ValueError, OverflowError, ZeroDivisionError) as error:
            raise HTTPException(422, "输入或运算结果不是可用的有限数值") from error
    elif name == "text_stats":
        text = arguments["text"]
        result = {"characters": len(text), "words": len(text.split()), "lines": len(text.splitlines())}
    elif name == "statistics":
        values = arguments["values"]
        if not values or not all(math.isfinite(value) for value in values):
            raise HTTPException(422, "请提供有限数值列表")
        result = {"count": len(values), "mean": statistics.mean(values), "median": statistics.median(values),
                  "stdev": statistics.stdev(values) if len(values) > 1 else 0, "min": min(values), "max": max(values)}
    elif name == "markdown_to_text":
        import html
        text = re.sub(r"!?\[([^\]]+)\]\([^)]+\)", r"\1", arguments["text"])
        result = html.unescape(re.sub(r"(?m)^\s*(?:#{1,6}\s+|>\s*)|[*_`~]+|<[^>]+>", "", text))
    elif name == "web_fetch":
        result = await web_fetch(arguments["url"])
    elif name == "web_search":
        result = await web_search(arguments["query"], arguments.get("max_results", 5))
    elif name == "knowledge_search":
        result = await asyncio.to_thread(store.search, "knowledge", arguments["query"], context.get("knowledge_ids") or [], None, arguments.get("limit", 5))
    elif name == "memory_read":
        result = await asyncio.to_thread(store.search, "memories", arguments.get("query", ""), None, context.get("memory_scope") or "global", arguments.get("limit", 5))
    elif name == "memory_save":
        result = await asyncio.to_thread(store.put, "memories", {"name": arguments.get("name") or arguments["content"][:80],
            "content": arguments["content"], "scope": context.get("memory_scope") or "global", "project_id": context.get("project_id")},
            None, {"job_id": context.get("job_id"), "source": "memory_save_tool"})
        result = {"id": result["id"], "name": result["name"], "scope": result["scope"], "saved": True}
    elif name == "project_read":
        from .projects import store as projects
        if not context.get("project_id") or not context.get("snapshot_id"):
            raise HTTPException(403, "未选择项目快照")
        snapshot = await asyncio.to_thread(projects.snapshot, context["project_id"], context["snapshot_id"])
        item = next((file for file in snapshot["files"] if file["path"] == arguments["path"]), None)
        if not item or not item.get("current"):
            raise HTTPException(404, "路径不在所选快照的当前源码中")
        current = item["current"]
        start = max(current["start_line"], arguments.get("start_line", current["start_line"]))
        end = min(current["end_line"], arguments.get("end_line", current["end_line"]))
        if end < start:
            raise HTTPException(422, "行范围不在快照覆盖范围内")
        result = {"path": item["path"], "start_line": start, "end_line": end,
                  "content": "\n".join(current["text"].splitlines()[start-current["start_line"]:end-current["start_line"]+1]),
                  "snapshot_id": snapshot["id"], "digest": snapshot["digest"]}
    elif name == "mcp_call":
        result = await mcp_call(arguments["server_id"], arguments["tool"], arguments.get("arguments", {}), context.get("mcp_server_ids") or [])
    return {"tool": name, "result": result}


@router.get("/overview")
async def overview():
    from .workbench import manager, status, TOOLS as root
    connection = await status()
    report_path = Path(__file__).resolve().parents[1] / "data/super-platform-validation.json"
    try:
        verified = json.loads(report_path.read_text()).get("checks", {}) if report_path.exists() else {}
    except (ValueError, OSError):
        verified = {}
    runtimes = [{"id": "gateway", "name": "已有账户模型网关", "installed": True, "available": connection["gateway"]["online"], "validated": connection["gateway"]["online"]},
                {"id": "camel", "name": "CAMEL", "version": "0.2.90", "installed": (root / ".venv-camel/bin/python").exists(), "validated": True},
                {"id": "crewai", "name": "CrewAI", "version": "1.15.27", "installed": (root / ".venv-crewai/bin/python").exists(), "validated": True}]
    features = [{"id": identifier, "name": name, "available": True, "validated": bool(verified.get(identifier)), "reason": reason} for identifier, name, reason in [
        ("knowledge", "知识检索与来源", "中文分词与 BM25 词项检索；不需要供应商 embedding key。"),
        ("memory", "持久记忆", "本机持久保存、记忆空间隔离与任务来源。"), ("mcp", "MCP", "官方 Python SDK；HTTP/SSE 与逐工具授权。"),
        ("flows", "任务流程", "原生 CAMEL Workforce 与 CrewAI Flow，运行状态保存在任务记录。"),
        ("sdk", "原生 SDK 工作室", "明确确认后在本机执行原生 SDK；额外服务与硬件按实际配置。"),
        ("migration", "数据迁移", "资源与历史增量导入，原始文件及当前数据保留。")]]
    return {"version": "2.0.0", "runtimes": runtimes, "counts": {kind: len(store.list(kind)) for kind in KINDS}, "features": features,
            "active_jobs": [item for item in manager.jobs.values() if item["status"] not in {"complete", "error", "cancelled", "interrupted"}],
            "retrieval": {"engine": "rank_bm25.BM25Plus", "tokenizer": "jieba", "semantic_embeddings": False}, "connection": connection}


@router.get("/tools")
async def tools_catalog():
    return [{**item, "available": True, "execution": "真实工具调用"} for item in TOOLS]


@router.post("/tools/{name}/invoke")
async def manual_tool(name: str, request: ToolRequest):
    context = request.model_dump(exclude={"arguments", "token"})
    validate_context(context)
    try:
        return await invoke_tool(name, request.arguments, context)
    except HTTPException:
        raise
    except Exception as error:
        from .workbench import public_error
        raise HTTPException(502, public_error(error)) from error


@router.post("/runtime/tools/{name}")
async def runtime_tool(name: str, request: ToolRequest):
    context = runtime_context(request.token)
    if name not in context["tools"]:
        raise HTTPException(403, "当前任务未授权此工具")
    try:
        result = await invoke_tool(name, request.arguments, context)
    except HTTPException:
        raise
    except Exception as error:
        from .workbench import public_error
        raise HTTPException(502, public_error(error)) from error
    from .workbench import manager
    job = manager.jobs.get(context["job_id"])
    if job:
        manager.emit(job, "tool_result", tool=name, output=json.dumps(result["result"], ensure_ascii=False)[:16000])
    return result


@router.post("/knowledge/search")
async def knowledge_search(request: SearchRequest):
    return await asyncio.to_thread(store.search, "knowledge", request.query, request.knowledge_ids, None, request.limit)


@router.post("/memories/search")
async def memories_search(request: SearchRequest):
    return await asyncio.to_thread(store.search, "memories", request.query, request.memory_ids, request.scope, request.limit)


@router.post("/knowledge/upload", status_code=201)
async def upload_knowledge(file: UploadFile = File(...)):
    name = Path(file.filename or "document.txt").name
    raw = await file.read(5_000_001)
    if len(raw) > 5_000_000:
        raise HTTPException(422, "上传文件不得超过 5 MB")
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise HTTPException(422, "请先解密 PDF")
        if len(reader.pages) > 200:
            raise HTTPException(422, "PDF 不得超过 200 页")
        content = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(raw))
        content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        content += "\n" + "\n".join("\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
    elif suffix in {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".yaml", ".yml", ".rst"}:
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise HTTPException(422, "文本文件需要 UTF-8 编码") from error
    else:
        raise HTTPException(422, "支持 PDF、DOCX、TXT、Markdown、CSV、JSON 和常见源码文本")
    from .projects import KNOWN_TOKEN, SECRET_CONTENT
    if KNOWN_TOKEN.search(content) or SECRET_CONTENT.search(content):
        raise HTTPException(422, "文档包含疑似凭据，请移除后上传")
    if not content.strip():
        raise HTTPException(422, "没有可提取文本；扫描 PDF 需要先进行 OCR")
    return await asyncio.to_thread(store.put, "knowledge", {"name": name, "content": content, "tags": [suffix.lstrip(".")]},
                                   None, {"source": "upload", "filename": name, "input_sha256": hashlib.sha256(raw).hexdigest()})


@router.post("/mcp/{identifier}/discover")
async def discover(identifier: str):
    try:
        return await mcp_discover(identifier)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(503, "MCP 服务器连接或初始化失败，请检查地址、权限和服务") from error


class MCPCallRequest(BaseModel):
    tool: str = Field(min_length=1, max_length=100)
    arguments: dict = Field(default_factory=dict)


@router.post("/mcp/{identifier}/call")
async def call_mcp(identifier: str, request: MCPCallRequest):
    try:
        return await mcp_call(identifier, request.tool, request.arguments)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(503, "MCP 调用失败，请检查工具参数与服务器状态") from error


@router.post("/flows/{identifier}/run", status_code=202)
async def run_flow(identifier: str, request: FlowRunRequest):
    from .workbench import JobRequest, create_job
    flow = store.get("flows", identifier)
    try:
        job = JobRequest(mode=flow["mode"], models=flow["models"], rounds=flow["rounds"], workflow=flow["workflow"], **request.model_dump(mode="json"))
    except ValueError as error:
        raise HTTPException(422, "请检查流程配置，并明确选择知识、记忆和项目上下文。") from error
    return await create_job(job)


@router.get("/sdk/templates")
async def sdk_templates():
    return [
        {"id": "camel-agent", "name": "CAMEL 原生 ChatAgent", "sdk": "camel", "source": 'from camel.agents import ChatAgent\nfrom camel.models import ModelFactory\nfrom camel.types import ModelPlatformType\nmodel = ModelFactory.create(model_platform=ModelPlatformType.OPENAI_COMPATIBLE_MODEL, model_type=MODELS[0], url=GATEWAY, api_key="local-oauth", model_config_dict={"max_tokens":1024})\nagent = ChatAgent("你是一名科研助手。", model=model)\nresponse = agent.step(PROMPT)\nRESULT = {"answer": response.msgs[0].content, "sdk":"camel"}\n'},
        {"id": "crewai-team", "name": "CrewAI 原生团队", "sdk": "crewai", "source": 'from crewai import Agent, Task, Crew, LLM\nllm = LLM(model="openai/"+MODELS[0], base_url=GATEWAY, api_key="local-oauth", max_tokens=1024)\nagent = Agent(role="分析员", goal="完成用户任务", backstory="重视证据", llm=llm)\ntask = Task(description=PROMPT, expected_output="简短中文结论", agent=agent)\noutput = Crew(agents=[agent], tasks=[task], verbose=False).kickoff()\nRESULT = {"answer": output.raw, "sdk":"crewai"}\n'},
    ]


@router.get("/data/export")
async def export_data(include_jobs: bool = True):
    from .workbench import manager
    data = {"format": "modeldesk-platform", "version": 2, "exported_at": timestamp(),
            "resources": {kind: store.list(kind) + store.list(kind, archived=True) for kind in KINDS},
            "jobs": [job for job in manager.jobs.values() if job["status"] in {"complete", "error", "cancelled", "interrupted"}] if include_jobs else []}
    content = json.dumps(data, ensure_ascii=False, indent=2)
    return Response(content=content, media_type="application/json", headers={"Content-Disposition": 'attachment; filename="ModelDesk-platform.json"', "Cache-Control": "no-store"})


class ImportRequest(BaseModel):
    data: dict
    conflict: Literal["skip", "rename"] = "skip"


def import_data(request):
    """Validate every record first; imports can never resurrect executable jobs."""
    data = request.data
    if len(json.dumps(data, ensure_ascii=False)) > 20_000_000:
        raise HTTPException(422, "导入数据超过 20 MB")
    resources, histories = data.get("resources"), data.get("jobs", [])
    if data.get("format") != "modeldesk-platform":
        # ChatHub exports provider credentials in sync/localStorage. Only prompts
        # and conversation text are accepted; those settings never migrate.
        if "local" not in data or "sync" not in data:
            raise HTTPException(422, "支持 ModelDesk 平台备份或 ChatHub 导出 JSON")
        local = data["local"]
        resources = {"prompts": [{"name": item.get("title", "导入提示词"), "content": item.get("prompt", ""), "category": "ChatHub 导入"} for item in local.get("prompts", [])]}
        histories = []
        for key, messages in local.items():
            if key.startswith("conversation:") and key.endswith(":messages") and isinstance(messages, list):
                text = "\n\n".join(str(item.get("role", "message")) + ": " + str(item.get("text", "")) for item in messages if isinstance(item, dict))
                if text.strip():
                    resources.setdefault("knowledge", []).append({"name": "ChatHub 历史 " + key.split(":")[2][:80], "content": text, "category": "导入历史"})
    if not isinstance(resources, dict) or set(resources) - KINDS or not isinstance(histories, list):
        raise HTTPException(422, "数据结构无效")
    validated = []
    for kind, items in resources.items():
        if not isinstance(items, list) or len(items) > 2000:
            raise HTTPException(422, "资源列表无效或超过上限")
        for item in items:
            if not isinstance(item, dict):
                raise HTTPException(422, "资源记录无效")
            body = {key: value for key, value in item.items() if key in Resource.model_fields}
            validated.append((kind, validate_resource(kind, body), item.get("id")))
    safe_jobs = []
    for job in histories:
        if not isinstance(job, dict) or job.get("status") not in {"complete", "error", "cancelled", "interrupted"}:
            raise HTTPException(422, "只能导入已结束的任务历史")
        try:
            identifier = str(uuid.UUID(job["id"]))
            conversation = str(uuid.UUID(job["conversation_id"]))
            if job["mode"] not in {"chat", "council", "camel", "crewai", "group", "review", "camel_workforce", "crewai_flow", "sdk_script"}:
                raise ValueError()
            models = job["models"]
            if not isinstance(models, list) or not 1 <= len(models) <= 4 or any(not isinstance(model, str) or not model or len(model) > 200 for model in models):
                raise ValueError()
            for key in ["created_at", "updated_at"]:
                if not isinstance(job[key], str) or datetime.fromisoformat(job[key].replace("Z", "+00:00")).tzinfo is None:
                    raise ValueError()
            if not isinstance(job.get("events"), list) or len(job["events"]) > 12003 or not isinstance(job.get("prompt"), str) or not job["prompt"].strip():
                raise ValueError()
            sequences = set()
            for event in job["events"]:
                if not isinstance(event, dict) or event.get("job_id") != identifier or not isinstance(event.get("type"), str) or not event["type"]:
                    raise ValueError()
                sequence = event.get("seq")
                if not isinstance(sequence, int) or isinstance(sequence, bool) or sequence < 1 or sequence in sequences:
                    raise ValueError()
                if datetime.fromisoformat(event["created_at"].replace("Z", "+00:00")).tzinfo is None:
                    raise ValueError()
                sequences.add(sequence)
            if job.get("result") is not None and not isinstance(job["result"], dict):
                raise ValueError()
            json.dumps(job, allow_nan=False)
            if any(key.startswith("_") or key in {"runtime", "token", "auth", "authorization"} for key in job):
                raise ValueError()
            safe_jobs.append({**job, "id": identifier, "conversation_id": conversation, "approved": False})
        except (KeyError, ValueError, TypeError) as error:
            raise HTTPException(422, "任务历史无效或包含私有运行字段") from error
    try:
        imported, skipped = store.import_resources(validated, request.conflict)
    except (ValueError, TypeError) as error:
        raise HTTPException(422, "资源标识无效，原数据保留") from error
    from .workbench import manager
    for job in safe_jobs:
        if job["id"] in manager.jobs:
            skipped += 1
            continue
        job["imported"] = True
        manager.jobs[job["id"]] = job
        manager.wakeups[job["id"]] = asyncio.Event()
        manager.save(job)
        imported += 1
    return {"imported": imported, "skipped": skipped, "existing_data_preserved": True, "credentials_imported": False}


@router.post("/data/import")
async def import_backup(request: ImportRequest):
    try:
        return import_data(request)
    except HTTPException:
        raise
    except (ValueError, TypeError) as error:
        raise HTTPException(422, "导入数据验证失败，检查记录字段") from error


@router.get("/{kind}")
async def list_resources(kind: str, archived: bool = False):
    items = await asyncio.to_thread(store.list, kind, archived)
    if kind == "knowledge":
        return [{**item, "chunks": len(item["chunks"]), "characters": len(item["content"])} for item in items]
    return items


@router.get("/{kind}/{identifier}")
async def get_resource(kind: str, identifier: str):
    return await asyncio.to_thread(store.get, kind, identifier)


@router.post("/{kind}", status_code=201)
async def create_resource(kind: str, request: Resource):
    return await asyncio.to_thread(store.put, kind, request.model_dump(mode="json"))


@router.put("/{kind}/{identifier}")
async def update_resource(kind: str, identifier: str, request: Resource):
    return await asyncio.to_thread(store.put, kind, request.model_dump(mode="json"), identifier)


@router.delete("/{kind}/{identifier}")
async def archive_resource(kind: str, identifier: str):
    return await asyncio.to_thread(store.archive, kind, identifier)


@router.post("/{kind}/{identifier}/restore")
async def restore_resource(kind: str, identifier: str):
    return await asyncio.to_thread(store.archive, kind, identifier, False)

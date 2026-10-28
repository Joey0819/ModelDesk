"""Immutable project assets, mature document extraction and job input manifests."""
import hashlib
import io
import json
import mimetypes
import os
from pathlib import Path, PurePosixPath
import stat
import threading
import uuid
import warnings
import zipfile

from fastapi import HTTPException
from .attachment_io import image_content_parts

ROOT = Path(__file__).resolve().parents[1]
MAX_FILE_BYTES = 20_000_000
MAX_BATCH_BYTES = 100_000_000
MAX_BATCH_FILES = 100
MAX_JOB_FILES = 30
MAX_TEXT = 100_000
MAX_JOB_TEXT = 16000
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".yaml", ".yml", ".rst", ".toml", ".ini", ".xml", ".sql", ".c", ".cpp", ".h", ".java", ".go", ".rs", ".sh"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def safe_relative(value):
    from .projects import excluded
    if not isinstance(value, str) or not value or len(value) > 1024 or any(char in value for char in ("\x00", "\n", "\r", "\\")):
        raise ValueError("附件相对路径无效")
    path = PurePosixPath(value)
    if path.is_absolute() or any(part in {"", ".", ".."} or part.startswith("-") for part in value.split("/")):
        raise ValueError("附件路径不能逃出项目目录")
    if excluded(value):
        raise ValueError("认证文件、凭据或缓存不能导入")
    return str(path)


def secret_text(text):
    from .projects import SECRET_CONTENT, UNQUOTED_SECRET, KNOWN_TOKEN, JSON_SECRET
    return any(pattern.search(text) for pattern in [SECRET_CONTENT, UNQUOTED_SECRET, KNOWN_TOKEN, JSON_SECRET])


def extract(raw, name):
    # Unknown formats are archived as unsupported, but cannot bypass credential
    # screening by changing the extension of a text credential file.
    if secret_text(raw.decode("utf-8", errors="ignore")):
        raise ValueError("文件包含疑似凭据，不能导入")
    suffix = Path(name).suffix.lower()
    mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
    result = {"mime_type": mime, "kind": "unsupported", "status": "unsupported", "reason": "此文件类型暂不支持读取", "truncated": False}
    if suffix in IMAGE_SUFFIXES:
        from PIL import Image, ImageOps
        if len(raw) > 10_000_000:
            raise ValueError("图片不得超过 10 MB")
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(raw)) as probe:
                    if probe.format not in {"PNG", "JPEG", "WEBP"} or getattr(probe, "n_frames", 1) != 1:
                        raise ValueError("只支持单帧 PNG、JPEG、WebP 图片")
                    if max(probe.size) > 8192 or probe.width * probe.height > 16_000_000:
                        raise ValueError("图片不得超过 8192 像素单边或 1600 万像素")
                    probe.verify()
                with Image.open(io.BytesIO(raw)) as decoded:
                    decoded.load()
                    converted = ImageOps.exif_transpose(decoded).convert("RGBA" if "A" in decoded.getbands() else "RGB")
                    clean = Image.frombytes(converted.mode, converted.size, converted.tobytes())
                    output = io.BytesIO()
                    fmt = "JPEG" if suffix in {".jpg", ".jpeg"} else "PNG"
                    clean.save(output, format=fmt, **({"quality": 95} if fmt == "JPEG" else {}))
                    payload = output.getvalue()
                    if len(payload) > 10_000_000:
                        raise ValueError("规范化后的图片超过 10 MB，请缩小图片")
                    return {"kind": "image", "status": "ready", "mime_type": "image/jpeg" if fmt == "JPEG" else "image/png", "width": clean.width, "height": clean.height, "truncated": False}, payload, ""
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("图片已损坏或不能安全解码") from error
    content = ""
    try:
        if suffix in TEXT_SUFFIXES:
            content = raw.decode("utf-8-sig")
            if "\x00" in content:
                raise ValueError("二进制内容不能作为文本导入")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 200:
                raise ValueError("PDF 需要解密且不得超过 200 页")
            pieces, length = [], 0
            for page in reader.pages:
                text = page.extract_text() or ""
                if len(text) > 2_000_000:
                    raise ValueError("PDF 单页文本超出安全读取范围")
                if secret_text(text):
                    raise ValueError("文件包含疑似凭据，不能导入")
                if length < MAX_TEXT:
                    pieces.append(text[:MAX_TEXT - length])
                length += len(text)
            content = "\n\n".join(pieces)
            result["truncated"] = length > MAX_TEXT
            if not content.strip():
                return {**result, "reason": "PDF 没有可提取文本；扫描件需要先进行 OCR"}, raw, ""
        elif suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 50_000_000:
                    raise ValueError("DOCX 解压体积超过安全限制")
            document = Document(io.BytesIO(raw))
            content = "\n".join(paragraph.text for paragraph in document.paragraphs)
            content += "\n" + "\n".join("\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            for section in document.sections:
                for container in [section.header, section.footer]:
                    content += "\n" + "\n".join(paragraph.text for paragraph in container.paragraphs)
                    content += "\n" + "\n".join("\t".join(cell.text for cell in row.cells) for table in container.tables for row in table.rows)
        else:
            return result, raw, ""
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("文件损坏、格式不匹配或文本不是 UTF-8") from error
    if secret_text(content):
        raise ValueError("文件包含疑似凭据，不能导入")
    result.update(kind="text", status="ready", reason=None, truncated=result["truncated"] or len(content) > MAX_TEXT)
    return result, raw, content[:MAX_TEXT]


class AssetStore:
    def __init__(self, directory=None):
        self.directory = Path(directory or ROOT / "data/assets")
        self.lock = threading.RLock()

    def record_path(self, identifier):
        try:
            return self.directory / str(uuid.UUID(identifier)) / "metadata.json"
        except (ValueError, TypeError) as error:
            raise HTTPException(422, "附件 ID 无效") from error

    def get(self, project_id, identifier):
        path = self.record_path(identifier)
        if not path.exists() or path.is_symlink() or path.parent.is_symlink() or self.directory.is_symlink():
            raise HTTPException(404, "附件不存在")
        item = json.loads(path.read_text())
        if item.get("project_id") != project_id:
            raise HTTPException(403, "附件不属于当前项目")
        return item

    def list(self, project_id):
        items = []
        for path in self.directory.glob("*/metadata.json"):
            try:
                item = json.loads(path.read_text())
                if item.get("project_id") == project_id:
                    items.append(item)
            except (ValueError, OSError):
                continue
        return sorted(items, key=lambda item: item["created_at"], reverse=True)

    def save(self, item):
        path = self.record_path(item["id"])
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(item, ensure_ascii=False, indent=2))
        temporary.chmod(0o600)
        temporary.replace(path)

    def add(self, project_id, name, raw, relative_path=None, shared=False):
        from .projects import store as projects
        projects.get(project_id)
        relative = safe_relative(relative_path or name)
        if len(raw) > MAX_FILE_BYTES:
            raise ValueError("每个文件不得超过 20 MB")
        details, payload, text = extract(raw, relative)
        identifier = str(uuid.uuid4())
        suffix = {"image/jpeg": ".jpg", "image/png": ".png"}.get(details["mime_type"], Path(relative).suffix.lower() or ".bin")
        directory = self.directory / identifier
        directory.mkdir(parents=True, mode=0o700)
        content = directory / ("content" + suffix)
        content.write_bytes(payload)
        content.chmod(0o600)
        (directory / "text.txt").write_text(text)
        (directory / "text.txt").chmod(0o600)
        from .workbench import now
        url = f"/api/workbench/projects/{project_id}/assets/{identifier}/content"
        item = {"id": identifier, "project_id": project_id, "name": PurePosixPath(relative).name, "relative_path": relative,
                "size": len(raw), "stored_size": len(payload), "sha256": sha(raw), "content_sha256": sha(payload),
                "text_sha256": sha(text.encode()), "text_preview": text[:1200], "shared": shared,
                "created_at": now(), "content_url": url, **details}
        if details["kind"] == "image":
            item["preview_url"] = url
        self.save(item)
        return item

    def content_path(self, item):
        directory = self.record_path(item["id"]).parent
        found = list(directory.glob("content.*"))
        if len(found) != 1 or found[0].is_symlink() or directory.is_symlink():
            raise HTTPException(409, "附件文件校验失败")
        raw = found[0].read_bytes()
        if len(raw) > MAX_FILE_BYTES or sha(raw) != item["content_sha256"]:
            raise HTTPException(409, "附件文件校验失败")
        return found[0]

    def update_shared(self, project_id, identifier, shared):
        with self.lock:
            item = self.get(project_id, identifier)
            item["shared"] = shared
            self.save(item)
            return item

    def remove(self, project_id, identifier):
        with self.lock:
            self.get(project_id, identifier)
            # Existing job copies retain the immutable bytes they selected.
            self.record_path(identifier).unlink()
        return {"removed": True, "job_copies_preserved": True}

    def validate(self, job):
        from .projects import store as projects
        if job.get("snapshot_id") and not job.get("project_id"):
            raise HTTPException(422, "快照必须关联项目")
        if job.get("project_id"):
            projects.get(job["project_id"])
        elif job.get("attachment_ids"):
            raise HTTPException(422, "请先为附件选择或创建项目")
        for identifier in job.get("attachment_ids", []):
            item = self.get(job["project_id"], identifier)
            if item["status"] != "ready":
                raise HTTPException(422, item.get("reason") or "附件暂不支持读取")
            self.content_path(item)

    def freeze(self, job, inputs_root):
        self.validate(job)
        if not job.get("project_id"):
            return None
        from .projects import store as projects
        project = projects.get(job["project_id"])
        selected = {identifier: "turn" for identifier in job.get("attachment_ids", [])}
        if job.get("include_project_assets", True):
            for item in self.list(job["project_id"]):
                if item.get("shared"):
                    selected.setdefault(item["id"], "project_shared")
        target = inputs_root / job["id"]
        target.mkdir(parents=True, exist_ok=True, mode=0o700)
        items, images, pieces, omitted, copies = [], [], [], [], []
        remaining, image_bytes, total_bytes = MAX_JOB_TEXT, 0, 0
        for identifier, source in selected.items():
            item = self.get(job["project_id"], identifier)
            entry = {key: item[key] for key in ["id", "project_id", "name", "relative_path", "mime_type", "kind", "status", "size", "sha256", "content_sha256", "truncated"]}
            entry.update(source=source)
            if len(items) >= MAX_JOB_FILES or item["status"] != "ready":
                omitted.append({"id": identifier, "reason": item.get("reason") or "达到每轮 30 个附件限制"})
                continue
            path = self.content_path(item)
            if total_bytes + item["stored_size"] > MAX_BATCH_BYTES:
                if source == "turn":
                    raise HTTPException(422, "所选附件总量超过 100 MB")
                omitted.append({"id": identifier, "reason": "达到每轮附件 100 MB 预算"})
                continue
            copied = target / (item["id"] + path.suffix)
            copied.write_bytes(path.read_bytes())
            copied.chmod(0o600)
            total_bytes += item["stored_size"]
            copies.append({"id": item["id"], "path": str(copied), "mime": item["mime_type"], "sha256": item["content_sha256"], "name": item["name"]})
            entry["content_url"] = f"/api/workbench/jobs/{job['id']}/attachments/{item['id']}/content"
            if item["kind"] == "image":
                if len(images) >= 8 or image_bytes + item["stored_size"] > 40_000_000:
                    if source == "turn":
                        raise HTTPException(422, "本轮图片超过 8 张或 40 MB，请减少所选图片")
                    omitted.append({"id": identifier, "reason": "达到本轮 8 张/40 MB 图片预算"})
                    continue
                images.append({"id": item["id"], "job_id": job["id"], "project_id": job["project_id"], "path": str(copied), "mime": item["mime_type"], "sha256": item["content_sha256"], "width": item["width"], "height": item["height"]})
                entry.update(width=item["width"], height=item["height"])
                entry["preview_url"] = entry["content_url"]
                image_bytes += item["stored_size"]
            else:
                text_path = path.parent / "text.txt"
                if text_path.is_symlink():
                    raise HTTPException(409, "附件文本校验失败")
                text = text_path.read_text()
                if sha(text.encode()) != item["text_sha256"]:
                    raise HTTPException(409, "附件文本校验失败")
                excerpt = text[:remaining]
                entry["text_preview"] = excerpt[:1200]
                entry["context_chars"] = len(excerpt)
                entry["truncated"] = entry["truncated"] or len(excerpt) < len(text)
                if not excerpt:
                    omitted.append({"id": identifier, "reason": "达到附件文本总预算"})
                    continue
                pieces.append(f"[附件 {item['relative_path']}；sha256={item['sha256']}]\n{excerpt}" + ("\n[此附件引用内容已截断]" if entry["truncated"] else ""))
                remaining -= len(excerpt)
            items.append(entry)
        profile = {"project_id": project["id"], "name": project["name"], "description": project.get("description", ""), "instructions": project.get("instructions", "")}
        material = "\n\n".join(pieces)
        manifest = {"project": profile, "items": items, "text_sha256": sha(material.encode()), "coverage": {"selected_count": len(selected), "included_count": len(items), "omitted": omitted, "truncated": bool(omitted) or any(item["truncated"] for item in items), "text_chars": MAX_JOB_TEXT - remaining, "image_count": len(images)}}
        manifest["digest"] = sha(json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode())
        data = {"manifest": manifest, "text": material, "images": images, "copies": copies}
        frozen = target / "context.json"
        frozen.write_text(json.dumps(data, ensure_ascii=False))
        frozen.chmod(0o600)
        return manifest

    def clone_frozen(self, source, source_root, job, inputs_root):
        """Copy a harness's immutable inputs; never reselect mutable shared assets."""
        import copy
        if source.get("project_id") != job.get("project_id"):
            raise HTTPException(403, "任务附件副本不能跨项目复用")
        if not source.get("attachment_manifest"):
            return None
        self.prepared(source, source_root)
        source_folder = source_root / source["id"]
        data = json.loads((source_folder / "context.json").read_text())
        manifest = copy.deepcopy(data["manifest"])
        target = inputs_root / job["id"]
        target.mkdir(parents=True, exist_ok=True, mode=0o700)
        copied_paths = {}
        for descriptor in data.get("copies", []):
            path = Path(descriptor["path"])
            if path.parent != source_folder:
                raise HTTPException(409, "冻结附件路径校验失败")
            raw = read_local(source_folder, path.name)
            if sha(raw) != descriptor["sha256"]:
                raise HTTPException(409, "冻结附件内容校验失败")
            copied = target / path.name
            copied.write_bytes(raw)
            copied.chmod(0o600)
            copied_paths[descriptor["id"]] = str(copied)
        for descriptor in [*data.get("copies", []), *data.get("images", [])]:
            descriptor["path"] = copied_paths[descriptor["id"]]
            if "job_id" in descriptor:
                descriptor["job_id"] = job["id"]
        for item in manifest["items"]:
            url = f"/api/workbench/jobs/{job['id']}/attachments/{item['id']}/content"
            item["content_url"] = url
            if item["kind"] == "image":
                item["preview_url"] = url
        manifest.pop("digest", None)
        manifest["digest"] = sha(json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode())
        data["manifest"] = manifest
        frozen = target / "context.json"
        frozen.write_text(json.dumps(data, ensure_ascii=False))
        frozen.chmod(0o600)
        return manifest

    def prepared(self, job, inputs_root):
        if not job.get("attachment_manifest"):
            return "", [], []
        path = inputs_root / job["id"] / "context.json"
        if path.is_symlink():
            raise HTTPException(409, "任务附件清单校验失败")
        data = json.loads(path.read_text())
        manifest = data["manifest"]
        digest = manifest.pop("digest")
        actual = sha(json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode())
        manifest["digest"] = digest
        if actual != digest or digest != job["attachment_manifest"]["digest"]:
            raise HTTPException(409, "任务附件清单校验失败")
        text = data["text"]
        if sha(text.encode()) != manifest["text_sha256"]:
            raise HTTPException(409, "任务附件文本校验失败")
        expected = {item["id"]: item for item in manifest["items"] if item["kind"] == "image"}
        if {image["id"] for image in data["images"]} != set(expected):
            raise HTTPException(409, "任务图片清单不完整")
        for image in data["images"]:
            if image["id"] not in expected or image["job_id"] != job["id"] or image["project_id"] != job["project_id"] or image["sha256"] != expected[image["id"]]["content_sha256"]:
                raise HTTPException(409, "任务图片清单校验失败")
        profile = manifest["project"]
        header = f"用户选择的项目：{profile['name']}\n项目说明：{profile['description']}\n项目要求：{profile['instructions']}"
        return header + ("\n\n<UNTRUSTED_ATTACHMENTS>\n" + text + "\n</UNTRUSTED_ATTACHMENTS>" if text else ""), manifest["items"], data["images"]

    def historical_images(self, job, inputs_root):
        """Reopen the effective historical copies, bound to their public digest.

        These are separate from context.json so the original selection and its
        immutable manifest never change when a later turn inherits an image.
        """
        expected_manifest = job.get("history_attachment_manifest")
        if not expected_manifest or not expected_manifest.get("items"):
            return [], []
        path = inputs_root / job["id"] / "history.json"
        if path.is_symlink() or path.parent.is_symlink():
            raise HTTPException(409, "历史图片清单校验失败")
        data = json.loads(path.read_text())
        manifest = data["manifest"]
        digest = manifest["digest"]
        unsigned = {key: value for key, value in manifest.items() if key != "digest"}
        if manifest != expected_manifest or sha(json.dumps(unsigned, ensure_ascii=False, sort_keys=True).encode()) != digest:
            raise HTTPException(409, "历史图片清单校验失败")
        expected = {item["id"]: item for item in manifest["items"]}
        if {image["id"] for image in data["images"]} != set(expected):
            raise HTTPException(409, "历史图片清单不完整")
        for image in data["images"]:
            if image["job_id"] != job["id"] or image["project_id"] != job["project_id"] or image["sha256"] != expected[image["id"]]["content_sha256"]:
                raise HTTPException(409, "历史图片副本校验失败")
        return manifest["items"], data["images"]

    def inherit_images(self, job, history, jobs, images, inputs_root):
        from .attachment_io import verified_image_bytes
        inherited, omitted, inherited_descriptors = [], [], []
        known = {item["id"] for item in images}
        total = sum(len(verified_image_bytes(item)) for item in images)
        for reference in reversed(history["metadata"]["references"]):
            previous = jobs.get(reference["job_id"])
            if not previous or previous.get("project_id") != job.get("project_id"):
                continue
            path = inputs_root / previous["id"] / "context.json"
            if not path.exists():
                for item in previous.get("attachments", []):
                    if item.get("kind") == "image":
                        omitted.append({"id": item["id"], "source_job_id": previous["id"], "reason": "历史图片副本不可用，不能重新查看原图"})
                continue
            self.prepared(previous, inputs_root)
            data = json.loads(path.read_text())
            historical_refs, historical_images = self.historical_images(previous, inputs_root)
            source_items = {item["id"]: item for item in [*data["manifest"]["items"], *historical_refs]}
            for descriptor in [*data["images"], *historical_images]:
                if descriptor["id"] in known:
                    continue
                raw = verified_image_bytes(descriptor)
                if len(images) >= 8 or len(job.get("attachments", [])) + len(inherited) >= MAX_JOB_FILES or total + len(raw) > 40_000_000:
                    omitted.append({"id": descriptor["id"], "source_job_id": previous["id"], "reason": "历史图片超出本轮附件预算"})
                    continue
                target = inputs_root / job["id"]
                target.mkdir(parents=True, exist_ok=True, mode=0o700)
                copied = target / Path(descriptor["path"]).name
                copied.write_bytes(raw)
                copied.chmod(0o600)
                copied_descriptor = {**descriptor, "job_id": job["id"], "path": str(copied)}
                images.append(copied_descriptor)
                inherited_descriptors.append(copied_descriptor)
                source = source_items[descriptor["id"]]
                url = f"/api/workbench/jobs/{job['id']}/attachments/{descriptor['id']}/content"
                inherited.append({**source, "source": "history", "source_job_id": source.get("source_job_id") or previous["id"], "content_url": url, "preview_url": url})
                total += len(raw)
                known.add(descriptor["id"])
        manifest = {"items": inherited, "coverage": {"included_count": len(inherited), "omitted": omitted, "truncated": bool(omitted)}}
        manifest["digest"] = sha(json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode())
        if inherited:
            target = inputs_root / job["id"]
            target.mkdir(parents=True, exist_ok=True, mode=0o700)
            path = target / "history.json"
            path.write_text(json.dumps({"manifest": manifest, "images": inherited_descriptors}, ensure_ascii=False))
            path.chmod(0o600)
        return manifest


def folder_files(folder):
    from .projects import excluded
    root = Path(folder["path"])
    if not root.is_dir() or root.is_symlink():
        return {"folder": {**folder, "status": "missing"}, "files": [], "omitted": [], "coverage": {"files_seen": 0, "files_listed": 0, "omitted_count": 0, "truncated": False, "limits": {"max_files": 1000}}}
    files, omitted, seen = [], [], 0
    for base, dirs, names in os.walk(root, followlinks=False):
        allowed = []
        for name in dirs:
            path = Path(base) / name
            relative = str(path.relative_to(root))
            if path.is_symlink() or excluded(relative):
                omitted.append({"relative_path": relative, "reason": "符号链接、缓存或认证目录已排除"})
            else:
                allowed.append(name)
        dirs[:] = sorted(allowed)
        for name in sorted(names):
            seen += 1
            path = Path(base) / name
            relative = str(path.relative_to(root))
            try:
                safe_relative(relative)
                if path.is_symlink() or not path.is_file():
                    raise ValueError("符号链接或非普通文件不能导入")
                size = path.stat().st_size
                if size > MAX_FILE_BYTES:
                    raise ValueError("文件超过 20 MB")
                suffix = path.suffix.lower()
                files.append({"relative_path": relative, "name": name, "size": size, "mime_type": mimetypes.guess_type(name)[0] or "application/octet-stream", "kind": "image" if suffix in IMAGE_SUFFIXES else "text" if suffix in TEXT_SUFFIXES | {".pdf", ".docx"} else "unsupported"})
            except (ValueError, OSError) as error:
                omitted.append({"relative_path": relative, "reason": str(error)})
            if seen >= 2000 or len(files) >= 1000:
                return {"folder": {**folder, "status": "available"}, "files": files, "omitted": omitted[:1000], "coverage": {"files_seen": seen, "files_listed": len(files), "omitted_count": len(omitted), "truncated": True, "limits": {"max_files": 1000}}}
    return {"folder": {**folder, "status": "available"}, "files": files, "omitted": omitted, "coverage": {"files_seen": seen, "files_listed": len(files), "omitted_count": len(omitted), "truncated": False, "limits": {"max_files": 1000}}}


def read_local(root, relative):
    relative = safe_relative(relative)
    root_fd = os.open(root, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    descriptor = root_fd
    try:
        parts = PurePosixPath(relative).parts
        for part in parts[:-1]:
            next_fd = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=descriptor)
            if descriptor != root_fd:
                os.close(descriptor)
            descriptor = next_fd
        fd = os.open(parts[-1], os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=descriptor)
        with os.fdopen(fd, "rb") as stream:
            if not stat.S_ISREG(os.fstat(stream.fileno()).st_mode):
                raise ValueError("只能导入普通文件")
            raw = stream.read(MAX_FILE_BYTES + 1)
        if len(raw) > MAX_FILE_BYTES:
            raise ValueError("文件超过 20 MB")
        return raw
    finally:
        if descriptor != root_fd:
            os.close(descriptor)
        os.close(root_fd)


store = AssetStore()
